"""
Beggars Map - Discovery Archive xlsx renderer.

Export/report artifact ONLY. Reads a JSON payload (written by
build-master-ledger.mjs) and writes a formatted .xlsx. Never touches
Supabase, never imports, never modifies the WIP xlsx, workbench-state.json
or excel-import-state.json. Deliberately zero business logic — every
decision about batch_id/reviewed/approved/archived/published/current_state
lives in build-master-ledger.mjs, so there is exactly one place to audit
that logic. This mirrors build-reverification-workbook.py's own existing
precedent for "render a formatted xlsx report from JSON state" exactly.

Two payload shapes, selected by payload["mode"]:
  "batch-export"  -> payload["batch_id"], payload["rows"] (raw WIP xlsx rows
                     for that batch — every WIP xlsx column, unchanged)
  "master-ledger" -> payload["rows"] (the ledger's own derived column set)

Usage: python build-archive-xlsx.py <payload.json> <dest.xlsx>
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATE_FILLS = {
    "PUBLISHED": PatternFill("solid", fgColor="C6EFCE"),
    "NOT_YET_PUBLISHED": PatternFill("solid", fgColor="FFEB9C"),
    "ARCHIVED": PatternFill("solid", fgColor="DDEBF7"),
    "APPROVED": PatternFill("solid", fgColor="E2EFDA"),
    "REVIEWED": PatternFill("solid", fgColor="F2F2F2"),
    "DISCOVERED": PatternFill("solid", fgColor="FFFFFF"),
}

# (header, json_key, width, wrap)
BATCH_EXPORT_COLUMNS = [
    ("place_id", "place_id", 30, False),
    ("Name", "name", 28, True),
    ("Address", "formatted_address", 36, True),
    ("Latitude", "latitude", 11, False),
    ("Longitude", "longitude", 11, False),
    ("Phone", "phone", 14, False),
    ("Number Valid", "Number Valid", 12, False),
    ("Menu List Under 100", "Menu List Under 100", 14, False),
    ("Menu Details/Notes", "Menu Details/Notes", 40, True),
    ("Excel Row", "_excelRow", 9, False),
]

MASTER_LEDGER_COLUMNS = [
    ("place_id", "place_id", 30, False),
    ("Name", "name", 26, True),
    ("Batch ID", "batch_id", 12, False),
    ("Current State", "current_state", 18, False),
    ("Reviewed", "reviewed", 9, False),
    ("Approved", "approved", 9, False),
    ("Export Status", "export_status", 13, False),
    ("Archived", "archived", 9, False),
    ("Archived At", "archived_at", 20, False),
    ("Workbench Purged", "workbench_purged", 15, False),
    ("Listing ID", "listing_id", 30, False),
    ("Published", "published", 10, False),
    ("Publication Date", "publication_date", 20, False),
    ("Photo Count (local)", "photo_count_local", 12, False),
    ("Photo Count (published)", "photo_count_published", 14, False),
    ("Photos Published", "photos_published", 12, False),
    ("Photos Archive-Only", "photos_archive_only", 14, False),
    ("Photos Archive Path", "photos_archive_path", 34, True),
    ("Batch Export Path", "batch_export_path", 40, True),
    ("Local Test Listing ID", "local_test_listing_id", 30, False),
    ("Known Hold", "known_hold", 40, True),
    ("Number Valid", "number_valid", 12, False),
    ("Menu List Under 100", "menu_list_under_100", 14, False),
    ("Menu Details/Notes", "menu_details_notes", 36, True),
    ("Excel Row", "excel_row", 9, False),
    ("Notes", "notes", 36, True),
]


def write_json(payload) -> None:
    sys.stdout.buffer.write(json.dumps(payload, ensure_ascii=False).encode("utf-8"))


def render(columns, rows, sheet_title):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    for col_idx, (header, _key, width, _wrap) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 22

    state_col = next((i for i, (_h, k, _w, _wr) in enumerate(columns, start=1) if k == "current_state"), None)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (_header, key, _width, wrap) in enumerate(columns, start=1):
            value = row.get(key, "")
            if isinstance(value, bool):
                value = "Yes" if value else "No"
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = BORDER
            if wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if state_col:
            fill = STATE_FILLS.get(row.get("current_state"))
            if fill:
                ws.cell(row=row_idx, column=state_col).fill = fill

    last_row = len(rows) + 1
    ws.freeze_panes = "C2"
    if last_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"

    return wb


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: build-archive-xlsx.py <payload.json> <dest.xlsx>", file=sys.stderr)
        return 2

    payload_path, dest_path = Path(sys.argv[1]), Path(sys.argv[2])
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    mode = payload.get("mode")

    if mode == "batch-export":
        wb = render(BATCH_EXPORT_COLUMNS, payload["rows"], f"Batch {payload['batch_id']}")
    elif mode == "master-ledger":
        wb = render(MASTER_LEDGER_COLUMNS, payload["rows"], "Master Discovery Ledger")
    else:
        print(f"Unknown payload mode: {mode!r}", file=sys.stderr)
        return 2

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest_path)
    write_json({"written": str(dest_path), "rows": len(payload["rows"])})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
